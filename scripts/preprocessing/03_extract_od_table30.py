import openpyxl
import pandas as pd
import csv

# Load the Excel file
excel_file = 'data/raw/od2023/Tabelas_Site_OD2023_REV_190225.xlsx'
wb = openpyxl.load_workbook(excel_file, data_only=True)
ws = wb['Tab30']

# Read the headers (zone numbers) from row 8
headers = []
for col_idx in range(2, 600):  # Start from column B (2)
    cell_value = ws.cell(8, col_idx).value
    if cell_value is None:
        break
    headers.append(int(cell_value) if isinstance(cell_value, (int, float)) else cell_value)

print(f"Found {len(headers)} destination zones")
print(f"First 10 zones: {headers[:10]}")
print(f"Last 10 zones: {headers[-10:]}")

# Read the zone origin data starting from row 10
zones = []
data_matrix = []

for row_idx in range(10, 1000):
    zone_id = ws.cell(row_idx, 1).value
    if zone_id is None:
        break
    zones.append(int(zone_id) if isinstance(zone_id, (int, float)) else zone_id)
    
    # Read the row data
    row_data = []
    for col_idx in range(2, len(headers) + 2):
        cell_value = ws.cell(row_idx, col_idx).value
        # Handle string representations of numbers
        if isinstance(cell_value, str):
            try:
                cell_value = float(cell_value)
            except:
                cell_value = 0
        if cell_value is None:
            cell_value = 0
        row_data.append(cell_value)
    
    data_matrix.append(row_data)

print(f"\nFound {len(zones)} origin zones")
print(f"Origin zones: {zones[:10]}...{zones[-5:]}")

# Convert to OD format: origin, destination, trips
od_data = []
for origin_idx, origin_zone in enumerate(zones):
    for dest_idx, dest_zone in enumerate(headers):
        trips = data_matrix[origin_idx][dest_idx]
        if trips != 0:  # Only include non-zero values
            od_data.append([origin_zone, dest_zone, trips])

print(f"\nTotal OD pairs (non-zero): {len(od_data)}")
print(f"First 10 OD pairs:")
for pair in od_data[:10]:
    print(pair)

# Write to CSV
output_file = 'data/processed/od_matrix_30.csv'
with open(output_file, 'w', newline='') as f:
    writer = csv.writer(f)
    writer.writerow(['origin', 'destination', 'trips'])
    for row in od_data:
        writer.writerow(row)

print(f"\nWritten {len(od_data)} rows to {output_file}")
